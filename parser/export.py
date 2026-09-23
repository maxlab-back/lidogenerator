"""Выгрузка лидов в Excel (вкладки по типам) и CSV.

Только openpyxl и stdlib: pandas/numpy сюда не тянем — на musl-дистрибутивах (Alpine),
нестандартных архитектурах и старых питонах их колёса не собираются, а из зависимостей
проекта они были самыми тяжёлыми.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .models import Lead, COLUMNS

SHEET_NAMES = {
    "podokonnik_pvh": "Подоконник ПВХ",
    "sandwich_otkos": "Сэндвич-панель откосы",
    "_maybe": "Под вопросом",
}
MAX_WIDTH_ROWS = 300      # по скольким строкам меряем ширину колонки


def _cell(v) -> object:
    """Значение для ячейки: списки и словари — строкой, None — пусто, числа как есть."""
    if v is None:
        return ""
    if isinstance(v, dict):
        return "; ".join(f"{k}: {', '.join(links)}" for k, links in v.items() if links)
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return v


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, head in enumerate(headers):
        seen = (str(r[i]) for r in rows[:MAX_WIDTH_ROWS])
        width = max([len(head), *(len(s) for s in seen)])
        ws.column_dimensions[get_column_letter(i + 1)].width = min(max(10, width + 2), 60)


def _save(wb: Workbook, target: str | Path | io.BytesIO) -> None:
    if "Sheet" in wb.sheetnames:            # пустая вкладка, созданная Workbook()
        del wb["Sheet"]
    if not wb.sheetnames:
        wb.create_sheet("Пусто").append(["нет данных"])
    wb.save(target)


def _write_csv(path: Path, headers: list[str], rows: list[list]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


# ----------------------------------------------------------------------------
# Нишевый режим (run.py): колонки из models.COLUMNS, сортировка по уверенности
# ----------------------------------------------------------------------------
def _niche_table(leads: list[Lead]) -> tuple[list[str], list[list]]:
    headers = [title for _, title in COLUMNS]
    rows = [l.to_row() for l in leads]
    rows.sort(key=lambda r: r.get("score") or 0, reverse=True)
    return headers, [[_cell(r.get(key, "")) for key, _ in COLUMNS] for r in rows]


# ----------------------------------------------------------------------------
# Универсальный режим (app.py): колонки выгрузки (ключ в JSON лида -> заголовок)
# ----------------------------------------------------------------------------
UNIVERSAL_COLUMNS = [
    ("name", "Название"),
    ("score", "Релевантность"),
    ("status", "Статус"),
    ("owner", "Менеджер"),
    ("country", "Страна"),
    ("region", "Регион"),
    ("city", "Город"),
    ("address", "Адрес"),
    ("phones", "Телефоны"),
    ("emails", "Почты"),
    ("socials", "Мессенджеры и соцсети"),
    ("website", "Сайт"),
    ("rubrics", "Рубрика"),
    ("rating", "Рейтинг"),
    ("ai_type", "Тип (ИИ)"),
    ("ai_reason", "Вывод ИИ"),
    ("inn", "ИНН / УНП"),
    ("legal_name", "Юрлицо"),
    ("legal_status", "Статус юрлица"),
    ("comment", "Комментарий"),
    ("source", "Источник"),
    ("source_url", "Ссылка"),
    ("query", "Запрос"),
    ("matched", "Совпало"),
]
_COUNTRY_NAMES = {"RU": "Россия", "BY": "Беларусь"}
_SOURCE_NAMES = {"maps": "Google Карты", "yandex_maps": "Яндекс Карты", "2gis": "2ГИС",
                 "search": "Google", "yandex": "Яндекс", "telegram": "Telegram"}


def _universal_table(rows: list[dict]) -> tuple[list[str], list[list]]:
    from .db import STATUSES

    headers = [title for _, title in UNIVERSAL_COLUMNS]
    out = []
    for r in rows:
        line = []
        for key, _ in UNIVERSAL_COLUMNS:
            v = r.get(key, "")
            if key == "country":
                v = _COUNTRY_NAMES.get(v, v)
            elif key == "status":
                v = STATUSES.get(v, v)
            elif key == "source":
                v = "+".join(_SOURCE_NAMES.get(s, s) for s in str(v).split("+") if s)
            line.append(_cell(v))
        out.append(line)
    return headers, out


def rows_to_xlsx(rows: list[dict], target: str | Path | io.BytesIO, threshold: float) -> None:
    """Excel с вкладками «Целевые» (score >= threshold) и «Под вопросом»."""
    good = [r for r in rows if (r.get("score") or 0) >= threshold]
    rest = [r for r in rows if (r.get("score") or 0) < threshold]
    wb = Workbook()
    for sheet, part in (("Целевые", good), ("Под вопросом", rest)):
        if not part and sheet == "Под вопросом":
            continue
        _write_sheet(wb, sheet, *_universal_table(part))
    _save(wb, target)


def export(leads: list[Lead], cfg: dict) -> dict:
    out_cfg = cfg.get("output", {})
    out_dir = Path(cfg.get("output", {}).get("dir", "output"))
    out_dir.mkdir(parents=True, exist_ok=True)
    threshold = cfg.get("score_threshold", 0.45)
    include_maybe = out_cfg.get("include_maybe", True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")

    # группируем
    groups: dict[str, list[Lead]] = {}
    maybe: list[Lead] = []
    for l in leads:
        if not l.lead_type or l.score < threshold:
            maybe.append(l)
        else:
            groups.setdefault(l.lead_type, []).append(l)

    written = {"counts": {}, "files": []}

    if out_cfg.get("excel", True):
        xlsx = out_dir / f"leads_{stamp}.xlsx"
        wb = Workbook()
        for key, sheet in SHEET_NAMES.items():
            data = (maybe if include_maybe else []) if key == "_maybe" else groups.get(key, [])
            if not data and key != "_maybe":
                continue
            if key == "_maybe" and not include_maybe:
                continue
            _write_sheet(wb, sheet, *_niche_table(data))
            written["counts"][sheet] = len(data)
        _save(wb, xlsx)
        written["files"].append(str(xlsx))

    if out_cfg.get("csv", True):
        for key, data in groups.items():
            name = SHEET_NAMES.get(key, key)
            path = out_dir / f"{key}_{stamp}.csv"
            _write_csv(path, *_niche_table(data))
            written["files"].append(str(path))
            written["counts"].setdefault(name, len(data))
        if include_maybe and maybe:
            path = out_dir / f"maybe_{stamp}.csv"
            _write_csv(path, *_niche_table(maybe))
            written["files"].append(str(path))
            written["counts"].setdefault(SHEET_NAMES["_maybe"], len(maybe))

    return written
